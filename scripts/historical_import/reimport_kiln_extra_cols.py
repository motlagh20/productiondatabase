#!/usr/bin/env python3
# SAFE partial re-import: add the 4 previously-dropped kiln sensor columns
# (باتوم2 -> bottom.B2, دمای لوله باتوم -> bottom.pipe,
#  دمای لوله خشک کن -> dryer.pipe, دمای واگن 44 -> wagon44.temp).
# Only INSERTs the missing readings; does NOT touch existing kiln_temperature_readings
# rows, so the 381 kiln_temp_correction links (by reading_id) stay intact.
# Idempotent: skips a (push_id, zone_group, zone_reading) triple already present.
import openpyxl, psycopg2, re, io
from pathlib import Path

RD=Path("C:/Users/Mohammad/Nextcloud/Projects/Trae/ProductionDatabase/xls/real data")

# column header (as in Excel) -> (zone_group, zone_reading)
COLS={
    "باتوم 2":("bottom","B2"),
    "دمای لوله باتوم":("bottom","pipe"),
    "دمای لوله خشک کن":("dryer","pipe"),
    "دمای واگن 44":("wagon44","temp"),
}
def clean(v):
    if v is None: return ""
    if isinstance(v,float): return str(int(v)) if v.is_integer() else str(v)
    return str(v).replace("\n"," ").strip()

CONN=dict(host="localhost",port=5433,dbname="postgres",user="postgres",password="test")
conn=psycopg2.connect(**CONN); cur=conn.cursor()

added=0; skipped=0; pushes_seen=0
for f in sorted(RD.glob("Kiln-*.xls*")):
    if f.name.startswith("~$"): continue
    try:
        wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
        sn="Input" if "Input" in wb.sheetnames else wb.sheetnames[0]
        ws=wb[sn]; rows=list(ws.iter_rows(values_only=True))
        if not rows: wb.close(); continue
        # Excel has a title row then the real header row; find the header row
        # (the one containing a known column like 'تاریخ').
        hdr_idx=0
        for ri,r in enumerate(rows[:5]):
            if any(str(h).replace(chr(10)," ").strip()=="تاریخ" for h in r if h is not None):
                hdr_idx=ri; break
        hdr=[clean(h) for h in rows[hdr_idx]]
        idx={h:i for i,h in enumerate(hdr)}
        # only keep headers we care about AND that exist in this file
        want={h:(ZONE_MAP_get := COLS[h]) for h in COLS if h in idx}
        if not want:
            wb.close(); continue
        for r in rows[1:]:
            date_j=clean(r[idx["تاریخ"]]) if "تاریخ" in idx and idx["تاریخ"]<len(r) else ""
            if not date_j: continue
            hour=clean(r[idx["ساعت"]]) if "ساعت" in idx and idx["ساعت"]<len(r) else ""
            incoming=clean(r[idx["واگن ورودی"]]) if "واگن ورودی" in idx else ""
            nk=f"{date_j}|{hour}|{incoming}"
            cur.execute("SELECT id FROM kiln_pushes WHERE natural_key=%s",(nk,))
            row=cur.fetchone()
            if not row:
                wb.close(); continue
            pid=row[0]; pushes_seen+=1
            for h,(grp,rd_) in want.items():
                ci=idx[h]
                val=clean(r[ci]) if ci<len(r) else ""
                if not val: continue
                # idempotent: skip if already present
                cur.execute("SELECT 1 FROM kiln_temperature_readings WHERE push_id=%s AND zone_group=%s AND zone_reading=%s",(pid,grp,rd_))
                if cur.fetchone():
                    skipped+=1; continue
                try: v=float(val)
                except:
                    continue
                cur.execute("INSERT INTO kiln_temperature_readings(push_id,zone_group,zone_reading,value,source) VALUES(%s,%s,%s,%s,'historical')",(pid,grp,rd_,v))
                added+=1
        wb.close()
    except Exception as e:
        print(f"SKIP {f.name}: {e}")
conn.commit()
print(f"Partial re-import: pushes_seen={pushes_seen}, readings_added={added}, skipped_dup={skipped}")
cur.execute("SELECT count(*) FROM kiln_temperature_readings"); print("total readings now:",cur.fetchone()[0])
cur.close(); conn.close()
