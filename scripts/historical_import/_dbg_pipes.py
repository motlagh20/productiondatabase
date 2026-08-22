import openpyxl
from pathlib import Path
RD=Path('C:/Users/Mohammad/Nextcloud/Projects/Trae/ProductionDatabase/xls/real data')
f=sorted(RD.glob('Kiln-1398.xls*'))[0]
wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
sn='Input' if 'Input' in wb.sheetnames else wb.sheetnames[0]
ws=wb[sn]; rows=list(ws.iter_rows(values_only=True))
hdr_idx=0
for ri,r in enumerate(rows[:5]):
    if any(str(h).replace(chr(10),' ').strip()=='تاریخ' for h in r if h is not None):
        hdr_idx=ri; break
hdr=[str(h).replace(chr(10),' ').strip() for h in rows[hdr_idx]]
idx={h:i for i,h in enumerate(hdr)}
print('hdr_idx', hdr_idx)
print('idx has لوله باتوم:', 'دمای لوله باتوم' in idx)
print('all cols (repr) with لوله or باتوم or واگن44:')
for i,h in enumerate(hdr):
    if any(k in h for k in ['لوله','باتوم','واگن44']):
        print('  ', i, repr(h))
ci=idx.get('دمای لوله باتوم')
print('ci for pipe:', ci)
for r in rows[hdr_idx+1:hdr_idx+6]:
    print('  val:', repr(r[ci]) if ci is not None and ci<len(r) else 'NO')
wb.close()
