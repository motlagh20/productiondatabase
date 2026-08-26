#!/usr/bin/env python3
# APPLY: fix Kiln-Merged.xlsx dates. Replicates the workbook's own control logic
# using COMPUTED values (read with data_only=True), writes with a normal workbook.
#   F (date_changed?) = IF(H_time - H_prev_time < 0, "Date Chaned", "")
#   C (date_control)  = IF(F="Date Chaned", IF(B_date==B_prev_date,"error",""), "")
# "error" => date_jalali needs +1 Jalali day (proper month/year rollover via jdatetime).
# flag (AF) = "check it" => NEVER touched (owner reviews with physical docs).
import openpyxl, jdatetime

PATH = r"C:/Projects/ProductionDatabase/xls/consolidated/Kiln-Merged.xlsx"

def to_min(s):
    if s is None: return None
    s = str(s).strip()
    if ":" not in s: return None
    try:
        h, m = s.split(":"); return int(h) * 60 + int(m)
    except Exception:
        return None

def next_jalali(s):
    y, m, d = map(int, str(s).split("."))
    dt = jdatetime.date(y, m, d)
    nxt = dt + jdatetime.timedelta(days=1)
    return "%04d.%02d.%02d" % (nxt.year, nxt.month, nxt.day)

# Read computed values
rd = openpyxl.load_workbook(PATH, data_only=True)["Kiln-All"]
rhdr = [c.value for c in rd[1]]
Bi, Hi, Fi, Afi = rhdr.index("date_jalali"), rhdr.index("PushingTime_min"), rhdr.index("date_changed?"), rhdr.index("flag")

prev_t = None; prev_d = None
plan = []   # (row, old_date, new_date)
for r in range(2, rd.max_row + 1):
    d = str(rd.cell(r, Bi + 1).value).strip() if rd.cell(r, Bi + 1).value else ""
    t = to_min(rd.cell(r, Hi + 1).value)
    fl = rd.cell(r, Afi + 1).value
    is_check = str(fl).strip().lower() in ("check it", "checkit", "check_it") if fl is not None else False
    chan = (prev_t is not None and t is not None and t < prev_t)
    err = chan and (prev_d is not None and d == prev_d)
    if err and not is_check and d:
        plan.append((r, d, next_jalali(d)))
    prev_t, prev_d = t, d

print("تعداد ردیف آماده‌ی اصلاح:", len(plan))
for p in plan[:8]: print("  ردیف %d: %s -> %s" % p)

# Write (normal workbook so formulas are preserved)
wb = openpyxl.load_workbook(PATH)
ws = wb["Kiln-All"]
whdr = [c.value for c in ws[1]]
wBi, wFi = whdr.index("date_jalali"), whdr.index("date_changed?")
for r, old, new in plan:
    ws.cell(r, wBi + 1).value = new
    ws.cell(r, wFi + 1).value = "%s -> %s" % (old, new)
wb.save(PATH)
print("اعمال شد. فایل ذخیره شد.")
print("کل تغییرات:", len(plan))
