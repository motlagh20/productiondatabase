#!/usr/bin/env python3
# TARGETED add-only: insert bottom.pipe + dryer.pipe readings that were dropped
# by the original loader (sheet mis-detection). Safe: only INSERTs missing rows,
# idempotent by (push_id, zone_group, zone_reading). Does NOT touch existing rows,
# so kiln_temp_correction (by reading_id) stays valid.
import openpyxl, psycopg2
from pathlib import Path

RD=Path("C:/Users/Mohammad/Nextcloud/Projects/Trae/ProductionDatabase/xls/real data")
# Excel header (exact) -> (zone_group, zone_reading)
COLS={
    "دمای لوله باتوم":("bottom","pipe"),
    "دمای لوله خشک کن":("dryer","pipe"),
}
CONN=dict(host="localhost",port=5433,dbname="postgres",user="postgres",password="test")
conn=psycopg2.connect(**CONN); cur=conn.cursor()

added=0; skip=0; files=0
for f in sorted(RD.glob("Kiln-*.xls*")):
    if f.name.startswith("~$"): continue
    try:
        wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
        sn="Input" if "Input" in wb.sheetnames else wb.sheetnames[0]
        ws=wb[sn]; rows=list(ws.iter_rows(values_only=True))
        # header row = first row containing 'تاریخ'
        hdr_idx=0
        for ri,r in enumerate(rows[:5]):
            if any(str(h).replace(chr(10)," ").strip()=="تاریخ" for h in r if h is not None):
                hdr_idx=ri; break
        hdr=[str(h).replace(chr(10)," ").strip() for h in rows[hdr_idx]]
        idx={h:i for i,h in enumerate(hdr)}
        want={h:COLS[h] for h in COLS if h in idx}
        if not want:
            wb.close(); continue
        files+=1
        for r in rows[hdr_idx+1:]:
            date_j=str(r[idx["تاریخ"]]).replace(chr(10)," ").strip() if "تاریخ" in idx and idx["تاریخ"]<len(r) else ""
            if not date_j: continue
            hour=str(r[idx["ساعت"]]).replace(chr(10)," ").strip() if "ساعت" in idx and idx["ساعت"]<len(r) else ""
            incoming=str(r[idx["واگن ورودی"]]).replace(chr(10)," ").strip() if "واگن ورودی" in idx else ""
            nk=f"{date_j}|{hour}|{incoming}"
            cur.execute("SELECT id FROM kiln_pushes WHERE natural_key=%s",(nk,))
            row=cur.fetchone()
            if not row: continue
            pid=row[0]
            for h,(grp,rd_) in want.items():
                ci=idx[h]
                v=r[ci] if ci<len(r) else None
                if v in (None,""): continue
                try: v=float(v)
                except: continue
                cur.execute("SELECT 1 FROM kiln_temperature_readings WHERE push_id=%s AND zone_group=%s AND zone_reading=%s",(pid,grp,rd_))
                if cur.fetchone():
                    skip+=1; continue
                cur.execute("INSERT INTO kiln_temperature_readings(push_id,zone_group,zone_reading,value,source) VALUES(%s,%s,%s,%s,'historical')",(pid,grp,rd_,v))
                added+=1
        wb.close()
    except Exception as e:
        print(f"SKIP {f.name}: {e}")
conn.commit()
print(f"Targeted pipe import: files={files}, added={added}, skipped_dup={skip}")
cur.execute("SELECT zone_group,zone_reading,count(*) FROM kiln_temperature_readings WHERE zone_group IN ('bottom','dryer') AND zone_reading IN ('pipe') GROUP BY 1,2")
print("pipe rows now:", cur.fetchall())
cur.close(); conn.close()
