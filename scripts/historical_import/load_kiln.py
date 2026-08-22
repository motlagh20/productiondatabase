#!/usr/bin/env python3
# Load KILN fact data into PostgreSQL. Push row + 18-zone temp readings (row-oriented).
# input_type canon {خشت خام, سفال پخته}; >1200C -> review. Read-only source.
import openpyxl, psycopg2, re
from pathlib import Path
RD=Path("C:/Users/Mohammad/Nextcloud/Projects/Trae/ProductionDatabase/xls/real data")
def clean(v):
    if v is None: return ""
    if isinstance(v,float): return str(int(v)) if v.is_integer() else str(v)
    return str(v).replace("\n"," ").strip()
CONN=dict(host="localhost",port=5433,dbname="postgres",user="postgres",password="test")
conn=psycopg2.connect(**CONN); cur=conn.cursor()
cur.execute("SELECT canonical_code,description FROM products"); prod_map={d:c for c,d in cur.fetchall()}
cur.close(); conn.close()

# zone column -> (group, reading) mapping (config vocab, Appendix C 12)
ZONE_MAP={
 "temp_exhaust":("exhaust","exhaust"),"دمای اگزوز":("exhaust","exhaust"),
 "temp_preheat01":("preheat","01"),"پیش گرما1":("preheat","01"),"preheat01":("preheat","01"),
 "temp_preheat02":("preheat","02"),"پیش گرما2":("preheat","02"),"preheat02":("preheat","02"),
 "temp_termostat":("thermostat","thermostat"),"ترموستات":("thermostat","thermostat"),
 "temp_Zone00":("zone","00"),"زون0":("zone","00"),"zone00":("zone","00"),"temp_zone00":("zone","00"),
 "temp_Zone01":("zone","01"),"زون1":("zone","01"),"zone01":("zone","01"),
 "temp_Zone02":("zone","02"),"زون2":("zone","02"),"zone02":("zone","02"),
 "temp_Zone03":("zone","03"),"زون3":("zone","03"),"zone03":("zone","03"),
 "temp_Zone04":("zone","04"),"زون4":("zone","04"),"zone04":("zone","04"),
 "temp_Zone05":("zone","05"),"زون5":("zone","05"),"zone05":("zone","05"),
 "temp_Zone06":("zone","06"),"زون6":("zone","06"),"zone06":("zone","06"),
 "temp_Zone07":("zone","07"),"زون7":("zone","07"),"zone07":("zone","07"),
 "temp_rapid01":("rapid","01"),"رپید1":("rapid","01"),
 "temp_rapid02":("rapid","02"),"رپید2":("rapid","02"),
 "temp_bottomA":("bottom","A"),"باتومَ A":("bottom","A"),"باتوم A":("bottom","A"),"bottomA":("bottom","A"),
 "temp_bottom01":("bottom","01"),"باتوم1":("bottom","01"),"bottom01":("bottom","01"),
 "temp_bottomB":("bottom","B"),"باتومB":("bottom","B"),"bottomB":("bottom","B"),
 "temp_bottom02":("bottom","B2"),"باتوم2":("bottom","B2"),"bottom02":("bottom","B2"),
 "دمای لوله باتوم":("bottom","pipe"),"دمای لوله خشک کن":("dryer","pipe"),
 "دمای واگن 44":("wagon44","temp"),
 }

kiln_files=sorted(RD.glob("Kiln-*.xls*"))
loaded=0; temps=0; flags=0
conn=psycopg2.connect(**CONN); cur=conn.cursor()
for f in kiln_files:
    if f.name.startswith("~$"): continue
    try:
        wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
        sn="Input" if "Input" in wb.sheetnames else wb.sheetnames[0]
        ws=wb[sn]; rows=list(ws.iter_rows(values_only=True))
        if not rows: wb.close(); continue
        hdr=[clean(h) for h in rows[0]]
        idx={h:i for i,h in enumerate(hdr)}
        # find temp columns present
        temp_cols={i:(ZONE_MAP.get(h,("unknown",h)) ) for i,h in enumerate(hdr) if h in ZONE_MAP}
        for r in rows[1:]:
            date_j=clean(r[idx["تاریخ"]]) if "تاریخ" in idx and idx["تاریخ"]<len(r) else ""
            if not date_j: continue
            hour=clean(r[idx["ساعت"]]) if "ساعت" in idx and idx["ساعت"]<len(r) else ""
            shift=clean(r[idx["شیفت"]]) if "شیفت" in idx and idx["شیفت"]<len(r) else ""
            pdesc=clean(r[idx["نام محصول"]]) if "نام محصول" in idx and idx["نام محصول"]<len(r) else ""
            pcode=prod_map.get(pdesc)
            # input_type
            itype=""
            if "خام" in idx and idx["خام"]<len(r) and clean(r[idx["خام"]]): itype=clean(r[idx["خام"]])
            elif "incomingProduct_Type" in idx and idx["incomingProduct_Type"]<len(r): itype=clean(r[idx["incomingProduct_Type"]])
            if itype in ("خام",): itype="خشت خام"
            elif itype in ("شارژی",): itype="سفال پخته"
            incoming=clean(r[idx["IncomingCarID"]]) if "IncomingCarID" in idx and idx["IncomingCarID"]<len(r) else (clean(r[idx["واگن ورودی"]]) if "واگن ورودی" in idx else "")
            ptime=clean(r[idx["زمان پوشینگ"]]) if "زمان پوشینگ" in idx and idx["زمان پوشینگ"]<len(r) else ""
            nk=f"{date_j}|{hour}|{incoming}"
            review=[]
            try:
                w=int(float(incoming)) if incoming else 0
                if w>80: review.append(("incoming_car_id",w,"Warning","flag>80"))
            except: pass
            cur.execute("""INSERT INTO kiln_pushes(date_jalali,hour,shift,product_code,input_type,incoming_car_id,pushing_time_min,push_seq,source,natural_key)
                         VALUES(%s,%s,%s,%s,%s,%s,%s,%s,'historical',%s)
                         ON CONFLICT (natural_key) DO UPDATE SET product_code=EXCLUDED.product_code""",
                        (date_j or None,hour or None,shift or None,pcode or None,itype or None,incoming or None,ptime or None,None,nk))
            cur.execute("SELECT id FROM kiln_pushes WHERE natural_key=%s",(nk,)); pid=cur.fetchone()[0]
            loaded+=1
            for ci,(grp,rd_) in temp_cols.items():
                val=clean(r[ci]) if ci<len(r) else ""
                if not val: continue
                try: v=float(val)
                except:
                    review.append((f"kiln_temp.{grp}",val,"Invalid","non-numeric")); continue
                if grp!="unknown" and v>1200:
                    review.append((f"kiln_temp.{grp}",v,"Invalid","x10 typo; divide by 10"))
                cur.execute("INSERT INTO kiln_temperature_readings(push_id,zone_group,zone_reading,value,source) VALUES(%s,%s,%s,%s,'historical')",(pid,grp,rd_,v))
                temps+=1
            for field,val,cls,fix in review:
                cur.execute("INSERT INTO review_queue(table_name,natural_key,field_name,raw_value,issue_class,suggested_fix) VALUES('kiln_pushes',%s,%s,%s,%s,%s)",(nk,field,str(val),cls,fix))
                flags+=1
        wb.close()
    except Exception as e:
        print(f"SKIP {f.name}: {e}")
conn.commit()
print(f"Kiln: pushes={loaded}, temp_readings={temps}, flagged={flags}")
cur.execute("SELECT count(*) FROM kiln_pushes"); print("total kiln_pushes:",cur.fetchone()[0])
cur.execute("SELECT count(*) FROM kiln_temperature_readings"); print("total kiln_temperature_readings:",cur.fetchone()[0])
cur.close(); conn.close()
